"""Build the translator's Word file from the defect Chinese.
Usage: python -I sample/build_sample_docx.py
Reads sample/report_zh_defects.md and writes sample/report_zh.docx: headings in the Heading 2 style,
front matter and numbered paragraphs as Normal paragraphs, each one run (no paragraph numbers in the
text; the engine maps paragraphs by order). Document properties name a fictional translator, with no generator mark.
"""
import os, re, sys
from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "report_zh_defects.md")
OUT = os.path.join(HERE, "report_zh.docx")

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
props = doc.core_properties
props.author = "Translator"
props.last_modified_by = "Translator"
props.comments = ""
props.created = datetime(2026, 5, 25, 9, 0, tzinfo=timezone.utc)
props.modified = datetime(2026, 5, 25, 17, 30, tzinfo=timezone.utc)
props.revision = 3
props.title = "德拉斯物流有限公司 截至2026年3月31日第一季度未经审计中期财务报告（中文译本）"

count = 0
for line in open(SRC, encoding="utf-8").read().splitlines():
    if line.startswith("## "):
        doc.add_heading(line[3:], level=2)
        continue
    m = re.match(r"^\[(F\d+|\d+)\]\s?(.*)$", line)
    if m:
        doc.add_paragraph(m.group(2))
        count += 1

doc.save(OUT)
print(f"build_sample_docx: {count} paragraphs written to {OUT}")
